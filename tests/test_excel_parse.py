"""Round-trip: export → simulate founder edits in Excel → parse back into SQLite."""
from pathlib import Path
import openpyxl


def _seed():
    from src.db import init_db
    from src.crm.leads import insert_lead
    init_db()
    insert_lead({"business_name": "Aurelia", "email": "a@aurelia.example.in", "industry": "architects"})
    insert_lead({"business_name": "PulseFit", "email": "p@pulsefit.example.in", "industry": "gyms_fitness_coaches"})


def test_leads_edit_roundtrip():
    _seed()
    from src.excel.exporter import export_leads
    p = export_leads()
    wb = openpyxl.load_workbook(str(p))
    ws = wb["All Leads"]
    headers = [c.value for c in ws[1]]
    appr_idx = headers.index("approval (approve/reject/leave blank)") + 1
    notes_idx = headers.index("founder_notes") + 1
    ws.cell(row=2, column=appr_idx, value="reject")
    ws.cell(row=2, column=notes_idx, value="not a fit right now")
    ws.cell(row=3, column=appr_idx, value="approve")
    wb.save(str(p))

    from src.excel.importer import parse_leads_edit
    counts = parse_leads_edit(p)
    assert counts["approved"] >= 1
    assert counts["rejected"] >= 1


def test_followups_edit_marks_done():
    from src.db import init_db, cursor
    from src.crm.leads import insert_lead
    init_db()
    lid = insert_lead({"business_name": "Skyline", "email": "i@skyline.example.ae"})
    with cursor() as cur:
        cur.execute("UPDATE leads SET next_followup_at=date('now','-1 day') WHERE lead_id=?", (lid,))

    from src.excel.exporter import export_followups_due
    p = export_followups_due()
    wb = openpyxl.load_workbook(str(p))
    ws = wb["Followups Due"]
    headers = [c.value for c in ws[1]]
    done_idx = headers.index("mark_done (yes/no)") + 1
    ws.cell(row=2, column=done_idx, value="yes")
    wb.save(str(p))

    from src.excel.importer import parse_followups_edit
    counts = parse_followups_edit(p)
    assert counts["done"] == 1


def test_auto_route_dispatches_by_filename(tmp_path):
    _seed()
    from src.excel.exporter import export_drafts
    from src.email.drafts import generate_drafts_for_new_leads
    generate_drafts_for_new_leads()
    p = Path(export_drafts())
    from src.excel.importer import auto_route
    kind, counters = auto_route(p)
    assert kind == "drafts"
    assert isinstance(counters, dict)
