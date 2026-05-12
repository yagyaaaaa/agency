from pathlib import Path
import openpyxl


def _seed_some():
    from src.db import init_db
    from src.crm.leads import insert_lead
    init_db()
    insert_lead({"business_name": "Aurelia Architects", "industry": "architects",
                 "city_country": "Jaipur, India", "website": "aurelia.example.in",
                 "email": "studio@aurelia.example.in", "personalization_hook": "test hook",
                 "website_quality_score": 4, "decision_maker": "Ms Anya"})
    insert_lead({"business_name": "Skyline Builders", "industry": "builders",
                 "city_country": "Dubai, UAE", "email": "info@skyline.example.ae",
                 "website_quality_score": 5})


def test_export_leads_excel():
    _seed_some()
    from src.excel.exporter import export_leads
    p = export_leads()
    assert Path(p).exists()
    wb = openpyxl.load_workbook(str(p))
    assert "All Leads" in wb.sheetnames
    assert "India Leads" in wb.sheetnames
    assert "Dubai Leads" in wb.sheetnames
    ws = wb["All Leads"]
    assert ws.freeze_panes == "A2"
    # first row are headers
    headers = [c.value for c in ws[1]]
    assert "business_name" in headers


def test_export_drafts_excel():
    _seed_some()
    from src.email.drafts import generate_drafts_for_new_leads
    n = generate_drafts_for_new_leads()
    assert n >= 1
    from src.excel.exporter import export_drafts
    p = export_drafts()
    assert Path(p).exists()
    wb = openpyxl.load_workbook(str(p))
    ws = wb.active
    headers = [c.value for c in ws[1]]
    assert "approve (yes/no)" in headers


def test_export_all_returns_paths():
    _seed_some()
    from src.excel.exporter import export_all
    paths = export_all()
    for key in ("leads", "drafts", "followups", "profit", "pipeline"):
        assert key in paths
        assert Path(paths[key]).exists()


def test_excel_importer_approves_drafts():
    _seed_some()
    from src.email.drafts import generate_drafts_for_new_leads, list_drafts
    generate_drafts_for_new_leads()
    from src.excel.exporter import export_drafts
    p = export_drafts()

    wb = openpyxl.load_workbook(str(p))
    ws = wb.active
    headers = [c.value for c in ws[1]]
    appr_col = headers.index("approve (yes/no)") + 1
    sent_col = headers.index("send_manually (yes/no)") + 1
    ws.cell(row=2, column=appr_col, value="yes")
    ws.cell(row=2, column=sent_col, value="yes")
    wb.save(str(p))

    from src.excel.importer import parse_drafts_edit
    counters = parse_drafts_edit(p)
    assert counters["approved"] >= 1
    assert counters["sent_manually"] >= 1

    sent = [d for d in list_drafts(status=None) if d["status"] == "sent_manually"]
    assert sent
