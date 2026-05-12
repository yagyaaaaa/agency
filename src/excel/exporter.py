"""Daily Excel exports.

Files written under data/excel/:
  leads_today_YYYY-MM-DD.xlsx
  cold_email_drafts_YYYY-MM-DD.xlsx
  followups_due_YYYY-MM-DD.xlsx
  profit_report_YYYY-MM-DD.xlsx
  pipeline_YYYY-MM-DD.xlsx

Conventions:
- header row frozen
- auto filter enabled
- approx auto-sized columns
- status / approval / notes columns where relevant
"""
from __future__ import annotations

from datetime import date
from pathlib import Path
from typing import Iterable, Sequence

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet

from src.db import cursor
from src.utils.paths import excel_path, today_str

HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
WRAP = Alignment(wrap_text=True, vertical="top")


def _autosize(ws: Worksheet, headers: Sequence[str]) -> None:
    for i, h in enumerate(headers, start=1):
        col_letter = openpyxl.utils.get_column_letter(i)
        max_len = max(
            [len(str(h))] + [len(str(ws.cell(row=r, column=i).value or "")) for r in range(2, min(ws.max_row, 200) + 1)]
        )
        ws.column_dimensions[col_letter].width = min(max(12, max_len + 2), 60)


def _write_sheet(ws: Worksheet, headers: Sequence[str], rows: Iterable[Sequence]) -> None:
    ws.append(list(headers))
    for cell in ws[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for r in rows:
        ws.append(list(r))
    ws.freeze_panes = "A2"
    if ws.max_row >= 2:
        ws.auto_filter.ref = ws.dimensions
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = WRAP
    _autosize(ws, headers)


# -------------- leads_today --------------

LEAD_HEADERS = [
    "lead_id", "business_name", "industry", "city_country", "website", "email",
    "phone", "decision_maker", "lead_score", "recommended_package", "status",
    "personalization_hook", "automation_opportunity", "notes",
    "approval (approve/reject/leave blank)", "founder_notes",
    "created_at", "updated_at",
]


def _fetch_leads(extra_where: str = "", args: tuple = ()) -> list[dict]:
    q = "SELECT * FROM leads"
    if extra_where:
        q += " WHERE " + extra_where
    q += " ORDER BY lead_score DESC, created_at DESC"
    with cursor() as cur:
        return [dict(r) for r in cur.execute(q, args).fetchall()]


def _lead_row(r: dict) -> list:
    return [
        r.get("lead_id"), r.get("business_name"), r.get("industry"),
        r.get("city_country"), r.get("website"), r.get("email"),
        r.get("phone"), r.get("decision_maker"), r.get("lead_score"),
        r.get("recommended_package"), r.get("status"),
        r.get("personalization_hook"), r.get("automation_opportunity"),
        r.get("notes"), "", "", r.get("created_at"), r.get("updated_at"),
    ]


def export_leads(day: str | None = None) -> Path:
    day = day or today_str()
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    all_leads = _fetch_leads()
    top = sorted(all_leads, key=lambda x: (x.get("lead_score") or 0), reverse=True)[:25]
    india = [r for r in all_leads if "india" in (r.get("city_country") or "").lower()]
    dubai = [r for r in all_leads if any(k in (r.get("city_country") or "").lower()
             for k in ("dubai", "uae", "abu dhabi"))]
    rejected = [r for r in all_leads if r.get("status") == "rejected"]
    contacted = [r for r in all_leads if r.get("status") in ("contacted", "followup_due", "sent_manually")]

    for name, rows in [
        ("All Leads", all_leads),
        ("Top Leads", top),
        ("India Leads", india),
        ("Dubai Leads", dubai),
        ("Rejected", rejected),
        ("Already Contacted", contacted),
    ]:
        ws = wb.create_sheet(name)
        _write_sheet(ws, LEAD_HEADERS, [_lead_row(r) for r in rows])

    out = excel_path("leads_today", day)
    wb.save(out)
    return out


# -------------- cold_email_drafts --------------

DRAFT_HEADERS = [
    "draft_id", "lead_id", "business_name", "email", "subject", "email_body",
    "personalization_reason", "recommended_package", "status",
    "approve (yes/no)", "founder_edits_to_subject", "founder_edits_to_body",
    "send_manually (yes/no)", "followup_1_date", "followup_2_date",
    "created_at",
]


def export_drafts(day: str | None = None) -> Path:
    day = day or today_str()
    with cursor() as cur:
        drafts = [dict(r) for r in cur.execute(
            "SELECT * FROM outreach_messages WHERE status IN ('needs_approval','approved','rejected') "
            "ORDER BY created_at DESC"
        ).fetchall()]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Drafts"

    rows = [
        [
            d.get("draft_id"), d.get("lead_id"), d.get("business_name"),
            d.get("email"), d.get("subject"), d.get("email_body"),
            d.get("personalization_reason"), d.get("recommended_package"),
            d.get("status"), "", "", "", "",
            d.get("followup_1_date"), d.get("followup_2_date"),
            d.get("created_at"),
        ]
        for d in drafts
    ]
    _write_sheet(ws, DRAFT_HEADERS, rows)
    out = excel_path("cold_email_drafts", day)
    wb.save(out)
    return out


# -------------- followups_due --------------

FOLLOWUP_HEADERS = [
    "lead_id", "business_name", "email", "next_followup_at", "status",
    "last_contacted_at", "notes", "mark_done (yes/no)", "founder_notes",
]


def export_followups_due(day: str | None = None) -> Path:
    day = day or today_str()
    with cursor() as cur:
        tracked = [dict(r) for r in cur.execute(
            """
            SELECT f.followup_id, f.lead_id, f.due_date, f.type, f.notes AS followup_notes,
                   l.business_name, l.email, l.status, l.last_contacted_at, l.notes
            FROM followups f
            JOIN leads l ON l.lead_id = f.lead_id
            WHERE f.done = 0 AND date(f.due_date) <= date('now')
            ORDER BY f.due_date ASC
            """
        ).fetchall()]
        leads = [dict(r) for r in cur.execute(
            "SELECT * FROM leads WHERE next_followup_at IS NOT NULL "
            "AND date(next_followup_at) <= date('now') "
            "ORDER BY next_followup_at ASC"
        ).fetchall()]
        drafts = [dict(r) for r in cur.execute(
            "SELECT * FROM outreach_messages WHERE status='approved' AND followup_1_date<=date('now') "
            "ORDER BY followup_1_date ASC"
        ).fetchall()]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Followups Due"

    rows = [
        [f.get("lead_id"), f.get("business_name"), f.get("email"),
         f.get("due_date"), f.get("status"),
         f.get("last_contacted_at"), f.get("followup_notes") or f.get("notes"), "", ""]
        for f in tracked
    ] + [
        [l.get("lead_id"), l.get("business_name"), l.get("email"),
         l.get("next_followup_at"), l.get("status"),
         l.get("last_contacted_at"), l.get("notes"), "", ""]
        for l in leads
    ]
    _write_sheet(ws, FOLLOWUP_HEADERS, rows)

    ws2 = wb.create_sheet("Email Followups Due")
    _write_sheet(ws2, [
        "draft_id", "lead_id", "business_name", "email", "followup_1_date",
        "followup_1_body", "followup_2_date", "founder_action (sent/skip)",
    ], [
        [d.get("draft_id"), d.get("lead_id"), d.get("business_name"),
         d.get("email"), d.get("followup_1_date"), d.get("followup_1_body"),
         d.get("followup_2_date"), ""]
        for d in drafts
    ])

    out = excel_path("followups_due", day)
    wb.save(out)
    return out


# -------------- profit_report --------------

PROFIT_REV_HEADERS = [
    "transaction_id", "client", "package", "add_ons", "amount_booked",
    "amount_collected", "payment_method", "payment_status",
    "commission_due", "net_profit_estimate", "notes", "created_at",
]
PROFIT_EXP_HEADERS = [
    "expense_id", "category", "vendor", "amount", "payment_method",
    "recurring", "notes", "created_at",
]


def export_profit(day: str | None = None) -> Path:
    day = day or today_str()
    with cursor() as cur:
        rev = [dict(r) for r in cur.execute("SELECT * FROM revenue ORDER BY created_at DESC").fetchall()]
        exp = [dict(r) for r in cur.execute("SELECT * FROM expenses ORDER BY created_at DESC").fetchall()]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Revenue"
    _write_sheet(ws, PROFIT_REV_HEADERS, [
        [r.get(h) for h in PROFIT_REV_HEADERS] for r in rev
    ])

    ws2 = wb.create_sheet("Expenses")
    _write_sheet(ws2, PROFIT_EXP_HEADERS, [
        [e.get(h) for h in PROFIT_EXP_HEADERS] for e in exp
    ])

    ws3 = wb.create_sheet("Summary")
    from src.profit.tracker import compute_summary
    s = compute_summary()
    _write_sheet(ws3, ["metric", "value"], list(s.items()))

    out = excel_path("profit_report", day)
    wb.save(out)
    return out


# -------------- pipeline --------------

PIPELINE_HEADERS = [
    "lead_id", "business_name", "industry", "city_country",
    "lead_score", "recommended_package", "status",
    "last_contacted_at", "next_followup_at", "notes",
]


def export_pipeline(day: str | None = None) -> Path:
    day = day or today_str()
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    stages = ["new", "contacted", "replied", "proposal_sent", "won", "lost", "rejected"]
    with cursor() as cur:
        for stage in stages:
            rows = [dict(r) for r in cur.execute(
                "SELECT * FROM leads WHERE status=? ORDER BY lead_score DESC", (stage,)
            ).fetchall()]
            ws = wb.create_sheet(stage.title())
            _write_sheet(ws, PIPELINE_HEADERS, [
                [r.get(h) for h in PIPELINE_HEADERS] for r in rows
            ])
    out = excel_path("pipeline", day)
    wb.save(out)
    return out


def export_all(day: str | None = None) -> dict[str, Path]:
    day = day or today_str()
    return {
        "leads": export_leads(day),
        "drafts": export_drafts(day),
        "followups": export_followups_due(day),
        "profit": export_profit(day),
        "pipeline": export_pipeline(day),
    }
