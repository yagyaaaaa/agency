"""Parse edited Excel files (uploaded back via Telegram) and update SQLite.

Recognized files by sheet/columns:
  - leads_today_*.xlsx: column 'approval (approve/reject/leave blank)' + 'founder_notes' on any sheet
  - cold_email_drafts_*.xlsx: column 'approve (yes/no)' + 'send_manually (yes/no)' + edits
  - followups_due_*.xlsx: column 'mark_done (yes/no)' + 'founder_notes'
"""
from __future__ import annotations

from pathlib import Path

import openpyxl

from src.crm.leads import set_status
from src.db import cursor
from src.email.drafts import mark_sent_manually, set_draft_status
from src.utils.logger import get_logger

log = get_logger("excel.importer")


def _open(path: str | Path):
    wb = openpyxl.load_workbook(filename=str(path), data_only=True)
    return wb


def _header_index(ws) -> dict[str, int]:
    headers = next(ws.iter_rows(values_only=True))
    return {(h or "").strip().lower(): i for i, h in enumerate(headers)}


def _yes(v) -> bool:
    if v is None:
        return False
    return str(v).strip().lower() in {"yes", "y", "true", "1", "approve", "approved"}


def _no(v) -> bool:
    if v is None:
        return False
    return str(v).strip().lower() in {"no", "n", "false", "0", "reject", "rejected"}


def parse_leads_edit(path: str | Path) -> dict[str, int]:
    """Update lead statuses + notes from the 'approval' and 'founder_notes' columns."""
    wb = _open(path)
    counters = {"approved": 0, "rejected": 0, "notes_appended": 0}

    for ws in wb.worksheets:
        idx = _header_index(ws)
        id_col = idx.get("lead_id")
        appr = idx.get("approval (approve/reject/leave blank)")
        notes_col = idx.get("founder_notes")
        if id_col is None:
            continue
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or row[id_col] in (None, ""):
                continue
            try:
                lead_id = int(row[id_col])
            except (TypeError, ValueError):
                continue
            decision = row[appr] if appr is not None else None
            note = row[notes_col] if notes_col is not None and notes_col < len(row) else None

            if _yes(decision):
                set_status(lead_id, "approved", note=str(note) if note else None)
                counters["approved"] += 1
            elif _no(decision):
                set_status(lead_id, "rejected", note=str(note) if note else None)
                counters["rejected"] += 1
            elif note:
                with cursor() as cur:
                    cur.execute(
                        "UPDATE leads SET notes=COALESCE(notes||char(10),'')||?, updated_at=CURRENT_TIMESTAMP "
                        "WHERE lead_id=?",
                        (str(note), lead_id),
                    )
                counters["notes_appended"] += 1
    log.info("parse_leads_edit %s -> %s", path, counters)
    return counters


def parse_drafts_edit(path: str | Path) -> dict[str, int]:
    wb = _open(path)
    counters = {"approved": 0, "rejected": 0, "sent_manually": 0, "edited": 0}

    ws = wb.worksheets[0]
    idx = _header_index(ws)
    id_col = idx.get("draft_id")
    appr = idx.get("approve (yes/no)")
    sent_col = idx.get("send_manually (yes/no)")
    subj_edit = idx.get("founder_edits_to_subject")
    body_edit = idx.get("founder_edits_to_body")
    if id_col is None:
        return counters

    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or row[id_col] in (None, ""):
            continue
        try:
            draft_id = int(row[id_col])
        except (TypeError, ValueError):
            continue
        decision = row[appr] if appr is not None else None
        sent = row[sent_col] if sent_col is not None and sent_col < len(row) else None
        new_subj = row[subj_edit] if subj_edit is not None and subj_edit < len(row) else None
        new_body = row[body_edit] if body_edit is not None and body_edit < len(row) else None

        edited = False
        if new_subj or new_body:
            with cursor() as cur:
                fields = []
                args: list = []
                if new_subj:
                    fields.append("subject=?")
                    args.append(str(new_subj))
                if new_body:
                    fields.append("email_body=?")
                    args.append(str(new_body))
                if fields:
                    args.append(draft_id)
                    cur.execute(
                        f"UPDATE outreach_messages SET {', '.join(fields)}, updated_at=CURRENT_TIMESTAMP WHERE draft_id=?",
                        tuple(args),
                    )
                    edited = True
        if edited:
            counters["edited"] += 1

        if _yes(decision):
            set_draft_status(draft_id, "approved", approved=True)
            counters["approved"] += 1
        elif _no(decision):
            set_draft_status(draft_id, "rejected", approved=False)
            counters["rejected"] += 1

        if _yes(sent):
            mark_sent_manually(draft_id)
            counters["sent_manually"] += 1
    log.info("parse_drafts_edit %s -> %s", path, counters)
    return counters


def parse_followups_edit(path: str | Path) -> dict[str, int]:
    wb = _open(path)
    counters = {"done": 0, "notes_appended": 0}
    ws = wb.worksheets[0]
    idx = _header_index(ws)
    id_col = idx.get("lead_id")
    done_col = idx.get("mark_done (yes/no)")
    notes_col = idx.get("founder_notes")
    if id_col is None:
        return counters

    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or row[id_col] in (None, ""):
            continue
        try:
            lead_id = int(row[id_col])
        except (TypeError, ValueError):
            continue
        done = row[done_col] if done_col is not None and done_col < len(row) else None
        note = row[notes_col] if notes_col is not None and notes_col < len(row) else None
        if _yes(done):
            with cursor() as cur:
                cur.execute(
                    "UPDATE leads SET next_followup_at=NULL, last_contacted_at=CURRENT_TIMESTAMP, "
                    "updated_at=CURRENT_TIMESTAMP WHERE lead_id=?",
                    (lead_id,),
                )
            counters["done"] += 1
        if note:
            with cursor() as cur:
                cur.execute(
                    "UPDATE leads SET notes=COALESCE(notes||char(10),'')||?, updated_at=CURRENT_TIMESTAMP "
                    "WHERE lead_id=?",
                    (str(note), lead_id),
                )
            counters["notes_appended"] += 1
    return counters


def auto_route(path: str | Path) -> tuple[str, dict]:
    """Decide which parser to run based on filename. Returns (kind, counters)."""
    name = Path(path).name.lower()
    if name.startswith("leads_today"):
        return "leads", parse_leads_edit(path)
    if name.startswith("cold_email_drafts"):
        return "drafts", parse_drafts_edit(path)
    if name.startswith("followups_due"):
        return "followups", parse_followups_edit(path)
    return "unknown", {}
